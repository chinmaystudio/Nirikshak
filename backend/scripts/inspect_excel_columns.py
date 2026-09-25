import openpyxl

wb_path = r"E:\Nirikshak\NIRIKSHAK_India_Infrastructure_Audit_Workbook_Pune_2000_to_2026.xlsx"
wb = openpyxl.load_workbook(wb_path, read_only=True)

print("Sheet names:", wb.sheetnames)

for name in ['Projects', 'Pune Deep Research', 'Source Registry', 'Data Dictionary']:
    if name in wb.sheetnames:
        sheet = wb[name]
        headers = []
        for row in sheet.iter_rows(values_only=True):
            headers = [str(c).strip() if c is not None else '' for c in row]
            break
        print(f"\n--- Sheet: {name} ---")
        print("Header count:", len(headers))
        print("Headers:", headers)
