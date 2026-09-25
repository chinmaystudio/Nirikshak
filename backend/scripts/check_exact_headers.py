import openpyxl

wb_path = r"E:\Nirikshak\NIRIKSHAK_India_Infrastructure_Audit_Workbook_Pune_2000_to_2026.xlsx"
wb = openpyxl.load_workbook(wb_path, read_only=True)

sheet = wb['Projects']
for r_i, r in enumerate(sheet.iter_rows(values_only=True), 1):
    if r_i == 3:
        print("Projects row 3 headers:", list(r))
        break

sheet_src = wb['Source Registry']
for r_i, r in enumerate(sheet_src.iter_rows(values_only=True), 1):
    if r_i == 3:
        print("\nSource Registry row 3 headers:", list(r))
        break

sheet_pune = wb['Pune Deep Research']
for r_i, r in enumerate(sheet_pune.iter_rows(values_only=True), 1):
    if r_i == 1:
        print("\nPune Deep Research row 1 headers:", list(r))
        break
